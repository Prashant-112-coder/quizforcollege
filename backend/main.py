import io, os, json, re
import fitz
import httpx
from docx import Document
from pptx import Presentation
from fastapi import FastAPI, UploadFile, File, Form, HTTPException
from fastapi.middleware.cors import CORSMiddleware
from pydantic import BaseModel

app=FastAPI(title="QuizForge AI API",version="1.0.0")
app.add_middleware(CORSMiddleware,allow_origins=["*"],allow_credentials=True,allow_methods=["*"],allow_headers=["*"])

class TopicRequest(BaseModel):
    topic:str
    count:int=10
    difficulty:str="mixed"
    exam_type:str="General"
    mode:str="exam"

def clean(t): return re.sub(r"\n{3,}","\n\n",re.sub(r"[ \t]+"," ",t)).strip()
def extract_pdf(data):
    d=fitz.open(stream=data,filetype="pdf")
    return [{"location":f"Page {i+1}","text":clean(p.get_text())} for i,p in enumerate(d)]
def extract_docx(data):
    d=Document(io.BytesIO(data))
    return [{"location":"Document","text":clean("\n".join(p.text for p in d.paragraphs if p.text.strip()))}]
def extract_pptx(data):
    p=Presentation(io.BytesIO(data)); out=[]
    for i,s in enumerate(p.slides):
        parts=[sh.text for sh in s.shapes if hasattr(sh,"text") and sh.text.strip()]
        out.append({"location":f"Slide {i+1}","text":clean("\n".join(parts))})
    return out
async def extract_file(name,data):
    ext=name.lower().rsplit(".",1)[-1] if "." in name else ""
    if ext=="pdf": return extract_pdf(data)
    if ext=="docx": return extract_docx(data)
    if ext=="pptx": return extract_pptx(data)
    if ext in ("txt","md","csv"): return [{"location":"Document","text":clean(data.decode("utf-8","ignore"))}]
    raise HTTPException(400,"Supported formats: PDF, DOCX, PPTX, TXT, MD, CSV")
def context(pages): return "\n\n".join(f"[{x['location']}]\n{x['text']}" for x in pages if x["text"])[:90000]
async def ai_generate(material,req):
    key=os.getenv("OPENAI_API_KEY")
    if not key:
        raise HTTPException(503,"AI service is not configured. Add OPENAI_API_KEY to the backend environment.")

    prompt=f"""Generate {req['count']} high-quality multiple-choice questions from the supplied study material. Difficulty={req['difficulty']}; exam={req['exam_type']}; mode={req['mode']}. Return ONLY JSON with title and questions. Each question must have question, options (4 strings), answer (0-3), explanation, source. Answers and explanations must be grounded in the material; never invent facts. MATERIAL:\n{material}"""
    payload={
        "model":os.getenv("OPENAI_MODEL","gpt-4.1-mini"),
        "messages":[
            {"role":"system","content":"You generate grounded educational MCQs."},
            {"role":"user","content":prompt}
        ],
        "temperature":0.2,
        "response_format":{"type":"json_object"}
    }

    async with httpx.AsyncClient(timeout=httpx.Timeout(90.0, connect=15.0)) as c:
        for attempt in range(3):
            try:
                r=await c.post(
                    "https://api.openai.com/v1/chat/completions",
                    headers={"Authorization":f"Bearer {key}"},
                    json=payload
                )
                if r.status_code == 429:
                    try:
                        err=r.json().get("error",{})
                    except ValueError:
                        err={}
                    message=err.get("message","OpenAI rate limit or quota exceeded.")
                    if attempt < 2 and "quota" not in message.lower() and "billing" not in message.lower():
                        await __import__("asyncio").sleep(2 ** attempt)
                        continue
                    raise HTTPException(429,f"OpenAI request limit reached: {message}")
                if r.status_code in (401,403):
                    raise HTTPException(r.status_code,"OpenAI authentication failed. Check OPENAI_API_KEY and project permissions.")
                if r.status_code >= 500:
                    if attempt < 2:
                        await __import__("asyncio").sleep(2 ** attempt)
                        continue
                    raise HTTPException(502,"OpenAI is temporarily unavailable. Please try again shortly.")
                r.raise_for_status()

                body=r.json()
                content=body.get("choices",[{}])[0].get("message",{}).get("content")
                if not content:
                    raise HTTPException(502,"OpenAI returned an empty quiz response.")
                try:
                    result=json.loads(content)
                except json.JSONDecodeError:
                    raise HTTPException(502,"OpenAI returned invalid quiz JSON.")
                if not isinstance(result,dict) or not isinstance(result.get("questions"),list):
                    raise HTTPException(502,"OpenAI returned an invalid quiz structure.")
                return result
            except httpx.TimeoutException:
                if attempt < 2:
                    await __import__("asyncio").sleep(2 ** attempt)
                    continue
                raise HTTPException(504,"OpenAI request timed out. Please try again.")
            except httpx.RequestError as exc:
                if attempt < 2:
                    await __import__("asyncio").sleep(2 ** attempt)
                    continue
                raise HTTPException(502,f"Unable to reach OpenAI: {exc.__class__.__name__}")

def fallback(material,count,title):
    sentences=[s.strip() for s in re.split(r"(?<=[.!?])\s+",material) if len(s.strip())>50]
    qs=[]
    for s in sentences[:count]:
        qs.append({"question":"Which statement is directly supported by the provided study material?","options":[s[:140],"This statement is not present in the material.","The material contradicts this statement.","None of these."],"answer":0,"explanation":s[:400],"source":"Uploaded material"})
    return {"title":title,"questions":qs}
async def build(material,title,req):
    result=await ai_generate(material,req)
    return result or fallback(material,req["count"],title)
@app.get("/health")
def health(): return {"status":"ok","service":"quizforge-backend"}
@app.post("/api/extract")
async def extract(file:UploadFile=File(...)):
    pages=await extract_file(file.filename,await file.read()); return {"filename":file.filename,"pages":pages,"text":context(pages)}
@app.post("/api/generate")
async def generate(file:UploadFile|None=File(None),topic:str|None=Form(None),count:int=Form(10),difficulty:str=Form("mixed"),exam_type:str=Form("General"),mode:str=Form("exam")):
    if file:
        pages=await extract_file(file.filename,await file.read()); material=context(pages); title=file.filename
    elif topic:
        material=f"Topic requested by learner: {topic}. Generate questions using your general knowledge."; title=topic
    else: raise HTTPException(400,"Upload a document or enter a topic.")
    req={"count":max(1,min(count,50)),"difficulty":difficulty,"exam_type":exam_type,"mode":mode}
    result=await build(material,title,req); result["source_filename"]=file.filename if file else None; return result
@app.post("/api/topic-quiz")
async def topic_quiz(req:TopicRequest):
    return await generate(topic=req.topic,count=req.count,difficulty=req.difficulty,exam_type=req.exam_type,mode=req.mode)
